"""Read-only workbook inspection and comparison with the published commit."""
import argparse
import io
import json
import math
import subprocess
import unicodedata
from pathlib import Path
import openpyxl

p = argparse.ArgumentParser()
p.add_argument('--sheet', action='append')
args = p.parse_args()
for path in Path('.').glob('*.xlsx'):
    current = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    before = subprocess.check_output(['git', 'show', f'ce85961:{path.name}'])
    old = openpyxl.load_workbook(io.BytesIO(before), data_only=False)
    old_cache = openpyxl.load_workbook(io.BytesIO(before), data_only=True)
    changed, cache_changed = [], []
    for s in old.sheetnames:
        for row in current[s]:
            for cell in row:
                a, b = old[s][cell.coordinate].value, cell.value
                if a != b:
                    changed.append(f'{s}!{cell.coordinate}')
                a, b = old_cache[s][cell.coordinate].value, cached[s][cell.coordinate].value
                equal = math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-12) if isinstance(a, (int,float)) and isinstance(b,(int,float)) else a == b
                if not equal:
                    cache_changed.append(f'{s}!{cell.coordinate}')
    print(unicodedata.normalize('NFC',path.name), len(old.sheetnames), '→', len(current.sheetnames))
    print('Added:', [s for s in current.sheetnames if s not in old.sheetnames])
    print('Changed expressions:', changed)
    print('Changed cached values (tolerance 1e-12):', cache_changed)
    print('Sheets:', cached.sheetnames)
    for s in args.sheet or []:
        if s in cached.sheetnames:
            print(s)
            for i,row in enumerate(cached[s],1):
                print(i, json.dumps([c.value for c in row],ensure_ascii=False,default=str))
