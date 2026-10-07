"""Read original workbooks without changing them; rebuild the dashboard dataset."""
import argparse
import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]

def clean(value):
    if value is None:
        return None
    if isinstance(value, (int, float, bool)):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    value = unicodedata.normalize("NFC", str(value))
    return value.strip()

# Workbook cell meaning, from 0-1 개요 and each sheet's legend. Other fills (header colors,
# 1-6 "해당" green, grey "?") are not data states and stay unmarked.
STATUS_FILLS = {"FFF2CC": "flagged", "E2EFDA": "verified"}

def cell_status(cell):
    """A cell can carry both marks, e.g. a yellow assumption typed in blue (3-3 B2)."""
    fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None
    color = cell.font.color.rgb if cell.font and cell.font.color is not None else None
    marks = [STATUS_FILLS[fill[-6:].upper()]] if isinstance(fill, str) and fill[-6:].upper() in STATUS_FILLS else []
    return marks + (["input"] if isinstance(color, str) and color[-6:].upper() == "0000FF" else [])

def key(name):
    name = unicodedata.normalize("NFKD", name).casefold()
    return re.sub(r"[^a-z0-9가-힣]", "", name)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=ROOT)
    args = parser.parse_args()
    aliases = json.loads((ROOT / "data/aliases.json").read_text())
    config = json.loads((ROOT / "data/import-config.json").read_text())
    hooks = json.loads((ROOT / "data/marketing-hypotheses.json").read_text())
    publisher_aliases = json.loads((ROOT / "data/publisher-aliases.json").read_text())
    alias_lookup = {key(name): key(group[0]) for group in aliases for name in group}
    def identity(name):
        return alias_lookup.get(key(name), key(name))

    workbooks = {}
    documents = []
    formula_books = {}
    for path in sorted(args.input_dir.glob("*.xlsx")):
        filename = unicodedata.normalize("NFC", path.name)
        workbook = openpyxl.load_workbook(path, data_only=True)
        formula_books[filename] = openpyxl.load_workbook(path, data_only=False)
        sheets = {}
        for sheet in workbook.worksheets:
            rows = [[clean(v) for v in row] for row in sheet.values]
            sheets[sheet.title] = rows
            formulas = {cell.coordinate: cell.value for cells in formula_books[filename][sheet.title] for cell in cells if cell.data_type == 'f'}
            blanks = {cell.coordinate: 'intentionalBlank' if '""' in cell.value else 'uncachedFormula'
                      for cells in formula_books[filename][sheet.title] for cell in cells
                      if cell.data_type == 'f' and sheet[cell.coordinate].value is None}
            status = {c.coordinate: s for cells in sheet.iter_rows() for c in cells
                      if c.value is not None and (s := cell_status(c))}
            documents.append({"id": key(filename + sheet.title), "file": filename,
                              "sheet": sheet.title, "rows": rows, "formulas": formulas, "blankStates": blanks,
                              "cellStatus": status})
        workbooks[filename] = sheets
    if not workbooks:
        raise ValueError("No XLSX workbooks found")
    market = next(v for k, v in workbooks.items() if "시장조사" in k)
    strategy = next(v for k, v in workbooks.items() if "게임전략" in k)
    overview = {r[0]: r[1] for r in market["0-1 개요"] if r[0] and r[1]}
    # 3-4 G column is the analyst's own App Store -> Google Play name match; honour it before any record is added.
    for values in market["3-4 App Store 매출 Top100"][1:]:
        if isinstance(values[1], str) and isinstance(values[6], str) and values[6] not in ("-", "") and key(values[1]) != key(values[6]):
            alias_lookup[key(values[1])] = alias_lookup.get(key(values[6]), key(values[6]))
    observations = []
    games = {}
    # Rows whose workbook cells are yellow (estimate / reading / needs check), per market sheet.
    flagged_rows = {d["sheet"]: {int(re.sub(r"[A-Z]+", "", c)) for c, v in d["cellStatus"].items() if "flagged" in v}
                    for d in documents if d["file"] == "시장조사.xlsx"}
    def add(name, publisher, genre, subgenre, country, store, date, rank, chart, sheet, row, notes=None, url=None):
        if not isinstance(rank, int) or not 1 <= rank <= 100:
            return
        missing = name == "확인 불가"
        game_id = identity(name) if not missing else f"missing-{country}-{store}-{chart}-{rank}"
        if game_id not in games:
            games[game_id] = {"id": game_id, "name": name, "publisher": publisher,
                              "developer": None, "genre": genre, "subgenre": subgenre,
                              "gameplay": None, "monetization": None, "marketing": None,
                              "artStyle": None, "audience": None, "playerAge": None,
                              "downloads": None, "installRange": None, "revenue": None, "analysis": None,
                              "evidence": None, "missing": missing, "marketingHook": None}
        observations.append({"id": f"{country}-{store}-{chart}-{date}-{rank}", "gameId": game_id,
                             "name": name, "publisher": publisher, "genre": genre, "subgenre": subgenre,
                             "country": country, "store": store, "date": date, "year": int(date[:4]),
                             "rank": rank, "chart": chart, "notes": notes, "flagged": row in flagged_rows.get(sheet, set()),
                             "source": {"file": "시장조사.xlsx", "sheet": sheet, "row": row, "url": url}})
    for store, sheet, url in [("ios", "1-1 App Store Top100", "https://gamedropdaily.com/mobile/"),
                               ("android", "1-2 Google Play Top100", "https://www.appbrain.com/stats/google-play-rankings/top_free/game/us")]:
        date = overview[("App Store" if store == "ios" else "Google Play") + " 기준일"]
        for row, values in enumerate(market[sheet][1:], 2):
            rank, name, publisher, genre, subgenre, notes = values[:6]
            add(name, publisher, genre, subgenre, "미국", store, date, rank, "free", sheet, row, notes, url)
    for sheet in ["2-2 국가별 Top20", "2-7 동남아 분석"]:
        for row, values in enumerate(market[sheet], 1):
            if len(values) >= 8 and isinstance(values[3], int) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(values[2])):
                _, country, date, rank, name, publisher, genre, source = values[:8]
                urls = re.findall(r"https?://\S+", str(source))
                add(name, publisher, genre, None, country, "android", date, rank, "free", sheet, row,
                    values[8] if len(values) > 8 else None, urls[0] if urls else None)
    for row, values in enumerate(market["3-2 매출 Top100"][1:], 2):
        rank, name, publisher, genre, subgenre, notes = values[:6]
        add(name, publisher, genre, subgenre, "미국", "android", config["grossingDate"], rank, "grossing",
            "3-2 매출 Top100", row, notes, "https://www.appbrain.com/stats/google-play-rankings/top_grossing/game/us")
    for row, values in enumerate(market["3-4 App Store 매출 Top100"][1:], 2):
        rank, name, publisher, genre, subgenre, basis, _, _, notes = values[:9]
        add(name, publisher, genre, subgenre, "미국", "ios", config["appStoreGrossingDate"], rank, "grossing",
            "3-4 App Store 매출 Top100", row, notes or basis, "https://www.appbrain.com/stats/appstore-rankings/top_grossing/games/us")
    for sheet in ["1-4 성공 비결 해부", "1-5 성공 비결 확장"]:
        for row, values in enumerate(strategy[sheet][1:], 2):
            if not values[0]:
                continue
            expanded = "확장" in sheet
            name = values[1] if expanded else values[0]
            game = games.get(identity(name)) if isinstance(name, str) else None
            if game:
                offset = 4 if expanded else 1
                game.update({"gameplay": values[offset], "monetization": values[6] if expanded else values[4],
                             "marketing": values[7] if expanded else values[5], "artStyle": values[8] if expanded else values[6],
                             "evidence": values[9] if expanded else values[7],
                             "analysis": {"file": "게임전략_MVP_마케팅.xlsx", "sheet": sheet, "row": row, "kind": "analystReading", "originalText": values[offset]}})
    install_status = next(d["cellStatus"] for d in documents if d["file"] == "시장조사.xlsx" and d["sheet"] == "3-3 매출-다운로드 비교")
    for row, values in enumerate(market["3-3 매출-다운로드 비교"][5:], 6):
        game = games.get(identity(values[0])) if values[0] else None
        if game and isinstance(values[3], (int, float)):
            game["downloads"] = {"value": values[3], "kind": "AppBrain 추정 누적 설치", "date": config["installEstimateDate"],
                                 "store": "android", "country": None}
        if game and values[5]:
            urls = re.findall(r"https?://\S+", str(values[8] or ""))
            game["installRange"] = {"range": values[5], "released": values[6],
                                    "flagged": [c for c in "FG" if "flagged" in install_status.get(f"{c}{row}", [])],
                                    "source": {"file": "시장조사.xlsx", "sheet": "3-3 매출-다운로드 비교", "row": row, "url": urls[0] if urls else None}}
    for name, hook in hooks.items():
        game = games.get(identity(name))
        if game:
            game["marketingHook"] = hook
    # Optional repeated observations are additive. Never rewrite the original XLSX.
    for snapshot_path in config.get('snapshotFiles', []):
        for record in json.loads((ROOT / snapshot_path).read_text()):
            required = ['name', 'publisher', 'genre', 'country', 'store', 'date', 'rank', 'chart', 'source']
            if any(k not in record for k in required) or record['store'] not in ['ios', 'android'] or record['chart'] not in ['free', 'grossing']:
                raise ValueError(f'Invalid ranking snapshot: {snapshot_path}')
            if type(record['rank']) is not int or not 1 <= record['rank'] <= 100 or not all(k in record['source'] for k in ['file','sheet','row']):
                raise ValueError('Snapshot requires rank 1..100 and file/sheet/row provenance')
            if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', record['date']):
                raise ValueError('Snapshot date must use YYYY-MM-DD')
            datetime.strptime(record['date'], '%Y-%m-%d')
            provenance=record['source']
            provenance['file']=unicodedata.normalize('NFC',provenance['file'])
            document=next((d for d in documents if d['file']==provenance['file'] and d['sheet']==provenance['sheet']),None)
            if not document or type(provenance['row']) is not int or not 1<=provenance['row']<=len(document['rows']):
                raise ValueError('Snapshot provenance must point to an imported workbook sheet and row')
            add(record['name'], record['publisher'], record['genre'], record.get('subgenre'), record['country'], record['store'], record['date'], record['rank'], record['chart'], record['source']['sheet'], record['source']['row'], record.get('notes'), record['source'].get('url'))
            observations[-1]['source'] = record['source']
    if len({o['id'] for o in observations}) != len(observations):
        raise ValueError('Duplicate country/store/chart/date/rank observation')

    from import_research import normalize_research
    research = normalize_research(market, strategy, identity)
    for game in games.values():
        game['tags'] = []
        for field, patterns in {
            'monetization': [('보상형 광고', r'보상형|리워드 광고'), ('IAP', r'IAP|인앱|결제|과금'), ('광고 중심', r'광고 중심|광고 수익형'), ('패스', r'패스')],
            'gameplay': [('정렬', r'정렬'), ('매치3', r'매치.?3'), ('머지', r'머지|합치'), ('블록', r'블록'), ('PvP', r'PvP|대전'), ('수집', r'수집')]
        }.items():
            for label, pattern in patterns:
                if game[field] and re.search(pattern, game[field], re.I):
                    kind = 'estimate' if '추정' in game[field] else 'analystReading'
                    game['tags'].append({'label': label, 'field': field, 'originalText': game[field], 'source': {**game['analysis'], 'kind': kind, 'originalText': game[field]}, 'kind': kind})
    # Site menu and update log come from each workbook's own 0 목차 / 0-1 개요, so new sheets need no code change.
    toc = [{"file": name, "items": [{"group": r[0], "sheet": r[1], "description": r[2], "status": r[3], "note": r[4]}
                                    for r in sheets["0 목차"][4:] if r[0] and r[1]]}
           for name, sheets in workbooks.items()]
    changelog = {}
    for row, values in enumerate(market["0-1 개요"], 1):
        match = re.match(r"(\d+단계(?:-\d+)?) (출처|한계)(?: \((\d{4}-\d{2}-\d{2}) 조회\))?", str(values[0] or ""))
        if match:
            entry = changelog.setdefault(match[1], {"stage": match[1], "date": None, "sources": None, "limits": None, "sheets": [], "row": row})
            entry["sources" if match[2] == "출처" else "limits"] = values[1]
            entry["date"] = entry["date"] or match[3]
    for book in toc:
        for item in book["items"]:
            for stage in re.findall(r"\d+단계(?:-\d+)?", str(item["note"] or "")):
                changelog.setdefault(stage, {"stage": stage, "date": None, "sources": None, "limits": None, "sheets": [], "row": None})["sheets"].append(
                    {"file": book["file"], "sheet": item["sheet"], "note": item["note"]})
    stage_order = lambda e: [int(n) for n in re.findall(r"\d+", e["stage"])]
    dataset = {"schemaVersion": 3, "importedAt": datetime.now(timezone.utc).isoformat(),
               "games": list(games.values()), "observations": observations, "documents": documents,
               "files": [{"name": name, "sheets": len(sheets)} for name, sheets in workbooks.items()], "publisherAliases": publisher_aliases,
               "toc": toc, "changelog": sorted(changelog.values(), key=stage_order, reverse=True), **research}
    (ROOT / "data/market.json").write_text(json.dumps(dataset, ensure_ascii=False, indent=2))
    print(f"Imported {len(games)} games, {len(observations)} observations, {len(documents)} sheets")

if __name__ == "__main__":
    main()
